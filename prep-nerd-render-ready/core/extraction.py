import hashlib
import ipaddress
import io
import socket
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup


def _safe_url(url):
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise ValueError("Only public HTTP/HTTPS URLs are supported")
    try:
        for result in socket.getaddrinfo(parsed.hostname, None):
            ip = ipaddress.ip_address(result[4][0])
            if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
                raise ValueError("Private/internal URLs are blocked")
    except socket.gaierror as exc:
        raise ValueError("Could not resolve URL") from exc
    return url


def extract_url(url, redirects=0):
    if redirects > 3:
        raise ValueError("Too many redirects")
    response = requests.get(_safe_url(url), timeout=15,
                            headers={"User-Agent": "prep-Nerd.ai/1.0"},
                            allow_redirects=False)
    if response.is_redirect:
        location = urljoin(url, response.headers.get("location", ""))
        return extract_url(location, redirects + 1)
    response.raise_for_status()
    if len(response.content) > 5_000_000:
        raise ValueError("URL response is too large")
    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text(" ", strip=True), response.url


def extract_file(uploaded):
    name = uploaded.name.lower()
    data = uploaded.getvalue()
    if len(data) > 20_000_000:
        raise ValueError("File exceeds 20 MB limit")
    if name.endswith((".txt", ".md", ".markdown")):
        return data.decode("utf-8", errors="replace")
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if name.endswith(".docx"):
        from docx import Document
        document = Document(io.BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if name.endswith(".pptx"):
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(data))
        output = []
        for index, slide in enumerate(presentation.slides, 1):
            output.append(f"[Slide {index}]")
            output.extend(shape.text for shape in slide.shapes if hasattr(shape, "text"))
        return "\n".join(output)
    if name.endswith(".ppt"):
        raise ValueError("Legacy .ppt is not supported in this environment; please upload .pptx")
    raise ValueError("Unsupported file type")


def content_hash(text):
    return hashlib.sha256(text.encode()).hexdigest()
